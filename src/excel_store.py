"""Excel 紀錄：一直往下加、被打開時重試。防護 4（公式注入）在這支檔案。"""
from __future__ import annotations

import json
import logging
import os
import time
from pathlib import Path
from typing import Callable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

HEADERS = [
    "日期時間", "上傳者", "模式類別", "查詢內容/品名",
    "主要分析結果", "備註與法規提醒", "照片存放位置",
]
WIDTHS = [20, 24, 12, 28, 60, 40, 48]

log = logging.getLogger(__name__)


def prepare_row(values: Sequence[object]) -> list[str]:
    row = ["" if v is None else str(v) for v in values]
    before = row

    # ==== 防護 4：Excel 文字開頭是 = + - @ 時轉成純文字 開始 ====
    # 前面加一個單引號，Excel 就會把整格當一般文字，不會當公式執行。
    # 所有文字格都要處理，包含上傳者（Discord 暱稱可以自己改成 =HYPERLINK(...)）。
    row = [("'" + v) if v.startswith(("=", "+", "-", "@", "\t", "\r")) else v for v in row]
    # ==== 防護 4 結束 ====

    changed = sum(a != b for a, b in zip(before, row))
    if changed:
        log.warning("防護 4：%d 個儲存格開頭是 = + - @，已轉成純文字", changed)
    return row


class ExcelStore:
    def __init__(
        self,
        path: Path,
        retries: int = 6,
        delay: float = 5.0,
        sleep: Callable[[float], None] = time.sleep,
        warn: Callable[[str], None] = log.warning,
    ) -> None:
        self.path = Path(path)
        self.pending_path = self.path.with_name("pending.jsonl")
        self.retries, self.delay, self.sleep, self.log = retries, delay, sleep, warn

    def append(self, values: Sequence[object]) -> bool:
        """新的一列加在最後。寫入成功回 True；Excel 被占用等原因寫不進去回 False（資料暫存，不遺失）。"""
        row = prepare_row(values)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        for attempt in range(1, self.retries + 1):
            try:
                backlog = self._load_pending()
                last_row = self._write(backlog + [row])
                self.pending_path.unlink(missing_ok=True)
                log.info("Excel 已寫入第 %d 列（%s）%s", last_row, self.path.name,
                         f"，其中補寫先前暫存的 {len(backlog)} 筆" if backlog else "")
                return True
            except Exception as exc:  # noqa: BLE001  任何原因都不能讓程式當掉
                self.log(
                    f"⚠️ 無法寫入 {self.path.name}（{type(exc).__name__}）："
                    f"如果你正開著它，請先關掉。第 {attempt}/{self.retries} 次嘗試。"
                )
                if attempt < self.retries:
                    self.sleep(self.delay)
        self._save_pending(row)
        self.log(f"⚠️ 這筆紀錄先暫存在 {self.pending_path.name}，下次寫入成功時會一併補上。")
        return False

    def _write(self, rows: list[list[str]]) -> int:
        """寫入並回傳最後一列的列號。"""
        if self.path.exists():
            wb = load_workbook(self.path)  # 打開既有檔案；壞掉會丟錯，絕不會被覆蓋
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "紀錄"
            ws.append(HEADERS)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col, width in enumerate(WIDTHS, start=1):
                ws.column_dimensions[ws.cell(1, col).column_letter].width = width
            ws.freeze_panes = "A2"
        for row in rows:
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        tmp = self.path.with_name(self.path.name + ".tmp")
        last_row = ws.max_row
        wb.save(tmp)
        os.replace(tmp, self.path)  # 先存暫存檔再換掉；寫到一半當掉也不會毀掉原檔
        return last_row

    def _load_pending(self) -> list[list[str]]:
        if not self.pending_path.exists():
            return []
        rows = []
        for line in self.pending_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                rows.append(json.loads(line))
        return rows

    def _save_pending(self, row: list[str]) -> None:
        try:
            with open(self.pending_path, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(row, ensure_ascii=False) + "\n")
        except OSError as exc:
            self.log(f"⚠️ 連暫存檔也寫不進去，這筆紀錄遺失：{exc}")
